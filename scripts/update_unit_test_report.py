"""
Script to update docs/Report5_Unit Test.xlsx from test-results-duc.json.
Maintains 100% formulas, formatting, and english text.
"""

import os
import json
import re
import openpyxl

WORKSPACE_DIR = r"C:\Users\DUCK-KANCHOU\Project_STAVE\S.T.A.V.E"
JSON_PATH = os.path.join(WORKSPACE_DIR, "test-results-duc.json")
EXCEL_PATH = os.path.join(WORKSPACE_DIR, "docs", "Report5_Unit Test.xlsx")

# Defect ID mapping for 29 failed cases
DEFECT_MAP = {
    # BE-DRF (11 cases)
    ("BE-DRF-01", "UTCID10"): "DFID001",
    ("BE-DRF-02", "UTCID08"): "DFID002",
    ("BE-DRF-03", "UTCID09"): "DFID003",
    ("BE-DRF-03", "UTCID10"): "DFID004",
    ("BE-DRF-04", "UTCID04"): "DFID005",
    ("BE-DRF-04", "UTCID07"): "DFID006",
    ("BE-DRF-04", "UTCID08"): "DFID007",
    ("BE-DRF-04", "UTCID11"): "DFID008",
    ("BE-DRF-04", "UTCID12"): "DFID009",
    ("BE-DRF-04", "UTCID14"): "DFID010",
    ("BE-DRF-05", "UTCID12"): "DFID011",
    # SH-TYP (1 case)
    ("SH-TYP-02", "UTCID12"): "DFID012",
    # BE-CMT Collaborator permissions (11 cases: DFID013 - DFID023)
    ("BE-CMT-01", "UTCID11"): "DFID013",
    ("BE-CMT-02", "UTCID11"): "DFID014",
    ("BE-CMT-03", "UTCID08"): "DFID015",
    ("BE-CMT-04", "UTCID11"): "DFID016",
    ("BE-CMT-05", "UTCID07"): "DFID017",
    ("BE-CMT-06", "UTCID06"): "DFID018",
    ("BE-CMT-07", "UTCID09"): "DFID019",
    ("BE-CMT-08", "UTCID11"): "DFID020",
    ("BE-CMT-09", "UTCID07"): "DFID021",
    ("BE-CMT-10", "UTCID06"): "DFID022",
    ("BE-CMT-15", "UTCID05"): "DFID023",
    # BE-CMT-21 DTO Bounds (5 cases: DFID024 - DFID028)
    ("BE-CMT-21", "UTCID10"): "DFID024",
    ("BE-CMT-21", "UTCID14"): "DFID025",
    ("BE-CMT-21", "UTCID17"): "DFID026",
    ("BE-CMT-21", "UTCID19"): "DFID027",
    ("BE-CMT-21", "UTCID20"): "DFID028",
    # BE-CMT-23 DTO message (1 case: DFID029)
    ("BE-CMT-23", "UTCID03"): "DFID029",
}

TARGET_SHEETS = (
    [f"BE-DRF-{i:02d}" for i in range(1, 6)] +
    [f"SH-TYP-{i:02d}" for i in range(1, 4)] +
    [f"FE-LIB-{i:02d}" for i in range(15, 18)] +
    [f"BE-CMT-{i:02d}" for i in range(1, 25)]
)

EXECUTED_BY = "Cam Trung Duc"
EXECUTED_DATE = "24/10/2026"

def load_test_results():
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)

    json_tests = {}
    for tr in data["testResults"]:
        for ar in tr["assertionResults"]:
            fn = ar["fullName"]
            sm = re.search(r"((?:BE|FE|SH)-[A-Z]+-\d+)", fn)
            um = re.search(r"UTCID\s*(\d+)", fn)
            if sm and um:
                s = sm.group(1)
                u = f"UTCID{int(um.group(1)):02d}"
                json_tests.setdefault(s, {})[u] = ar["status"]
    return json_tests

def main():
    print("Loading test results from JSON...")
    json_tests = load_test_results()

    print(f"Loading workbook: {EXCEL_PATH} ...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False)

    stats_summary = {
        "sheets_updated": 0,
        "cases_total": 0,
        "cases_passed": 0,
        "cases_failed": 0,
        "defects_recorded": 0,
    }

    for sname in TARGET_SHEETS:
        if sname not in wb.sheetnames:
            print(f"WARNING: Sheet {sname} not found in workbook!")
            continue

        ws = wb[sname]
        stats_summary["sheets_updated"] += 1

        # 1. Update Executed By at L3 (cell 3, 12)
        ws["L3"].value = EXECUTED_BY

        # 2. Locate row numbers for UTCID, Passed/Failed, Executed Date, Defect ID, Note
        utcid_row = 9
        pf_row = None
        date_row = None
        defect_row = None
        note_row = None

        for r in range(1, ws.max_row + 1):
            v2 = str(ws.cell(r, 2).value or "").strip().lower()
            if v2 == "passed/failed":
                pf_row = r
            elif v2 == "executed date":
                date_row = r
            elif v2 == "defect id":
                defect_row = r
            elif v2 in ["note", "notes", "remark", "remarks"]:
                note_row = r

        if not pf_row or not date_row or not defect_row:
            raise ValueError(f"Could not find required rows in {sname}: pf={pf_row}, date={date_row}, defect={defect_row}")

        # 3. Process UTCID columns
        for c in range(6, ws.max_column + 1):
            cell_val = ws.cell(utcid_row, c).value
            if cell_val and str(cell_val).startswith("UTCID"):
                utcid = str(cell_val).strip()
                stats_summary["cases_total"] += 1

                st = json_tests.get(sname, {}).get(utcid)
                if st is None:
                    print(f"WARNING: No test result found for {sname} {utcid}!")
                    continue

                if st == "passed":
                    ws.cell(pf_row, c).value = "P"
                    ws.cell(defect_row, c).value = None
                    stats_summary["cases_passed"] += 1
                else:  # failed
                    ws.cell(pf_row, c).value = "F"
                    dfid = DEFECT_MAP.get((sname, utcid))
                    if dfid:
                        ws.cell(defect_row, c).value = dfid
                        stats_summary["defects_recorded"] += 1
                    else:
                        print(f"ERROR: No defect ID mapped for failed test {sname} {utcid}!")
                    stats_summary["cases_failed"] += 1

                # Executed Date
                ws.cell(date_row, c).value = EXECUTED_DATE

        # 4. Mandatory Technical Notes (100% English)
        if sname == "BE-DRF-05":
            # UTCID03 is column 8 (H)
            ws.cell(45, 8).value = "Allow 0 tracks (empty snapshot defaults to tracks: [] per team decision, does not violate validation)."
        elif sname == "SH-TYP-02":
            # UTCID12 is column 17 (Q)
            ws.cell(39, 17).value = "Duplicate UUID across 2 pitches returns true because the function only compares entry count after sanitize (expected fail per sheet specification)."
        elif sname == "FE-LIB-15":
            note_text = "Test formatMonthYear function (4 test cases) according to sheet FE-LIB-15 specification."
            ws.cell(26, 3).value = note_text
            ws.cell(26, 6).value = note_text
        elif sname == "FE-LIB-16":
            note_text = "Test formatDate function (6 test cases) according to sheet FE-LIB-16 specification."
            ws.cell(30, 3).value = note_text
            ws.cell(30, 6).value = note_text
        elif sname == "FE-LIB-17":
            note_text = "Test getInitials function (frontend/lib/format-name.ts) according to sheet FE-LIB-17 specification (truncateHash function does not exist in frontend source code)."
            ws.cell(31, 3).value = note_text
            ws.cell(31, 6).value = note_text

    print("Saving modified workbook...")
    wb.save(EXCEL_PATH)
    print("Workbook saved successfully!")
    print(f"Summary: {stats_summary}")

if __name__ == "__main__":
    main()
