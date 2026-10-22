import re, sys, zipfile, os, openpyxl
from xml.sax.saxutils import escape
from openpyxl.utils import get_column_letter

UNIT_CASES = {
    "BE-BRN-04": ["UTCID04"],
    "BE-BRN-08": ["UTCID05", "UTCID06"],
    "BE-BRN-09": ["UTCID10"],
    "BE-BRN-10": ["UTCID04"],
    "BE-BRN-12": ["UTCID07"],
}
PENDING = {"Branch Management": ["BRN-044", "BRN-045", "BRN-048"]}


def sheet_path(z, name):
    wbx = z.read("xl/workbook.xml").decode()
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid = re.search(r'<sheet name="%s"[^>]*r:id="([^"]+)"' % re.escape(escape(name, {'"': "&quot;"})), wbx).group(1)
    return "xl/" + re.search(r'<Relationship Id="%s"[^>]*Target="([^"]+)"' % rid, rels).group(1)


def set_cell(xml, ref, text):
    pat = re.compile(r'<c r="%s"( s="\d+")?(?: t="\w+")?\s*(?:/>|>.*?</c>)' % ref, re.S)
    m = pat.search(xml)
    if not m:
        raise SystemExit("cell %s not found" % ref)
    if text is None:
        cell = '<c r="%s"%s/>' % (ref, m.group(1) or "")
    else:
        cell = '<c r="%s"%s t="inlineStr"><is><t>%s</t></is></c>' % (ref, m.group(1) or "", escape(text))
    return xml[: m.start()] + cell + xml[m.end():]


def patch(path, ops):
    z = zipfile.ZipFile(path)
    changed = {}
    for sheet, cells in ops.items():
        p = sheet_path(z, sheet)
        xml = z.read(p).decode()
        for ref, val in cells:
            xml = set_cell(xml, ref, val)
        changed[p] = xml.encode()
    wbx = z.read("xl/workbook.xml").decode()
    if "fullCalcOnLoad" not in wbx:
        wbx = re.sub(r"<calcPr([^>]*)/>", r'<calcPr\1 fullCalcOnLoad="1"/>', wbx)
    changed["xl/workbook.xml"] = wbx.encode()
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w") as out:
        for info in z.infolist():
            out.writestr(info, changed.get(info.filename, z.read(info.filename)), compress_type=zipfile.ZIP_DEFLATED)
    z.close()
    os.replace(tmp, path)


def unit_ops(path):
    wb = openpyxl.load_workbook(path)
    ops = {}
    for s, cases in UNIT_CASES.items():
        ws = wb[s]
        hdr = pf = dt = None
        cols = {}
        for row in ws.iter_rows(max_row=14):
            for c in row:
                if isinstance(c.value, str) and re.fullmatch(r"UTCID\d+", c.value.strip()):
                    cols[c.value.strip()] = c.column
                    hdr = c.row
        for row in ws.iter_rows(min_row=hdr):
            for c in row[:3]:
                if isinstance(c.value, str):
                    v = c.value.strip()
                    if v.startswith("Passed/"): pf = c.row
                    if v == "Executed Date": dt = c.row
        for u in cases:
            L = get_column_letter(cols[u])
            ops.setdefault(s, []).extend([("%s%d" % (L, pf), "P"), ("%s%d" % (L, dt), "21/10/2026")])
    return ops


def report_ops(path):
    wb = openpyxl.load_workbook(path)
    ops = {}
    for s, ids in PENDING.items():
        ws = wb[s]
        for row in ws.iter_rows(min_row=11):
            if row[0].value in ids:
                ops.setdefault(s, []).extend([("F%d" % row[0].row, "Pending"), ("G%d" % row[0].row, None)])
    return ops


if __name__ == "__main__":
    kind, path = sys.argv[1], sys.argv[2]
    ops = unit_ops(path) if kind == "unit" else report_ops(path)
    if len(sys.argv) > 3 and sys.argv[3] == "dry":
        print(ops)
    else:
        patch(path, ops)
        print("patched", kind, sum(len(v) for v in ops.values()), "cells")
