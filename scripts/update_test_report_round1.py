"""
Script to update Round 1 (Integration Testing) results for Version Control (VC)
and Project Management (PRJ) in docs/Report5_Test Report.xlsx.
Maintains 100% formulas, styles, and data integrity.
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment

sys.stdout.reconfigure(encoding='utf-8')

# Paths
BASE_DIR = r"C:\Users\DUCK-KANCHOU\Project_STAVE\S.T.A.V.E"
DRIVE_SOURCE = r"C:\Users\DUCK-KANCHOU\Downloads\Report5_Test_Report_GoogleDrive.xlsx"
LOCAL_TARGET = os.path.join(BASE_DIR, "docs", "Report5_Test Report.xlsx")
ONEDRIVE_TARGET = r"C:\Users\DUCK-KANCHOU\OneDrive\文档\STAVE\Report5_Test Report.xlsx"

TESTER_DUC = "Ducctde180199"
TEST_DATE = "24/10/2026"

# 12 Failed cases for Version Control with their corresponding DFID
VC_DEFECT_MAP = {
    "VC-005": "DFID024",  # UC-42: Commit unavailable when draft has no changes (BR-40)
    "VC-013": "DFID013",  # UC-43: View commit history of active branch
    "VC-016": "DFID014",  # UC-43: Project with only initial commit history
    "VC-020": "DFID019",  # UC-44: Add tag to commit
    "VC-023": "DFID025",  # UC-44: Empty tag name rejected (BR-44)
    "VC-026": "DFID021",  # UC-44: Tagging commit with existing tag replacement (BR-44)
    "VC-029": "DFID023",  # UC-45: Compare commits on two different branches (BR-46)
    "VC-035": "DFID020",  # UC-46: Restore earlier version as new commit (BR-47)
    "VC-041": "DFID001",  # UC-116: Working draft saved automatically (BR-130)
    "VC-043": "DFID006",  # UC-116: Nothing to autosave after commit until edit
    "VC-049": "DFID007",  # UC-116: Autosave interval 30s-1m boundary (BR-130)
    "VC-050": "DFID003",  # UC-116: Draft recovery after unexpected end (BR-131)
}

# Backend API Use Cases for Project Management
PRJ_BACKEND_UCS = {"UC-18", "UC-20", "UC-21", "UC-23"}

def update_test_report():
    print(f"Loading workbook from: {DRIVE_SOURCE} ...")
    wb = openpyxl.load_workbook(DRIVE_SOURCE, data_only=False)

    font_tahoma_10 = Font(name="Tahoma", size=10, bold=False, color="000000")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # ==========================================
    # 1. Update Version Control Sheet
    # ==========================================
    ws_vc = wb["Version Control"]
    vc_stats = {"Passed": 0, "Failed": 0, "Pending": 0, "Total": 0}

    for r in range(12, ws_vc.max_row + 1):
        tc_id = ws_vc.cell(r, 1).value
        if not tc_id or not str(tc_id).startswith("VC-"):
            continue

        tc_str = str(tc_id).strip()
        vc_stats["Total"] += 1
        res = ws_vc.cell(r, 6).value  # Col F: Round 1 Result

        # Tester (Col H)
        ws_vc.cell(r, 8).value = TESTER_DUC
        ws_vc.cell(r, 8).font = font_tahoma_10
        ws_vc.cell(r, 8).alignment = align_left

        # Existing base tag from Note
        raw_note = str(ws_vc.cell(r, 15).value or "").strip()
        # Clean existing suffixes if any
        base_tag = raw_note.split(" (")[0].split(" - ")[0].strip()

        if res == "Passed":
            vc_stats["Passed"] += 1
            ws_vc.cell(r, 7).value = TEST_DATE
            ws_vc.cell(r, 7).font = font_tahoma_10
            ws_vc.cell(r, 7).alignment = align_left

            # Note
            ws_vc.cell(r, 15).value = f"{base_tag} (Automated - Supertest / Postman)"
            ws_vc.cell(r, 15).font = font_tahoma_10
            ws_vc.cell(r, 15).alignment = align_left

        elif res == "Failed":
            vc_stats["Failed"] += 1
            ws_vc.cell(r, 7).value = TEST_DATE
            ws_vc.cell(r, 7).font = font_tahoma_10
            ws_vc.cell(r, 7).alignment = align_left

            dfid = VC_DEFECT_MAP.get(tc_str, "DFID013")
            ws_vc.cell(r, 15).value = f"{base_tag} ({dfid})"
            ws_vc.cell(r, 15).font = font_tahoma_10
            ws_vc.cell(r, 15).alignment = align_left

        else:  # Pending
            vc_stats["Pending"] += 1
            ws_vc.cell(r, 7).value = None  # Pending has no test date
            if "(Manual)" in raw_note:
                ws_vc.cell(r, 15).value = f"{base_tag} (Manual)"
            else:
                ws_vc.cell(r, 15).value = base_tag
            ws_vc.cell(r, 15).font = font_tahoma_10
            ws_vc.cell(r, 15).alignment = align_left

    print(f"Version Control updated: {vc_stats}")

    # ==========================================
    # 2. Update Project Management Sheet
    # ==========================================
    ws_prj = wb["Project Management"]
    prj_stats = {"Passed": 0, "Failed": 0, "Pending": 0, "Total": 0}
    current_uc = None

    for r in range(11, ws_prj.max_row + 1):
        cell_a = ws_prj.cell(r, 1).value
        if cell_a and str(cell_a).startswith("UC-"):
            current_uc = str(cell_a).split()[0].strip()
            continue

        if not cell_a or not str(cell_a).startswith("PRJ-"):
            continue

        tc_str = str(cell_a).strip()
        prj_stats["Total"] += 1
        res = ws_prj.cell(r, 6).value  # Col F: Round 1 Result

        if res == "Passed":
            prj_stats["Passed"] += 1
        elif res == "Failed":
            prj_stats["Failed"] += 1
        else:
            prj_stats["Pending"] += 1

        # Check if backend API UC (Create project, Archive, Delete, Visibility)
        is_backend_api = current_uc in PRJ_BACKEND_UCS

        if is_backend_api:
            # Synchronize Tester and Test Date for backend API cases
            ws_prj.cell(r, 8).value = TESTER_DUC
            ws_prj.cell(r, 8).font = font_tahoma_10
            ws_prj.cell(r, 8).alignment = align_left

            if res in ["Passed", "Failed"]:
                ws_prj.cell(r, 7).value = TEST_DATE
            else:
                ws_prj.cell(r, 7).value = None
            ws_prj.cell(r, 7).font = font_tahoma_10
            ws_prj.cell(r, 7).alignment = align_left
        else:
            # UC-19 and UC-22 maintain Nguyen Van Tri
            ws_prj.cell(r, 8).value = "Nguyen Van Tri"
            ws_prj.cell(r, 8).font = font_tahoma_10
            ws_prj.cell(r, 8).alignment = align_left

            if res in ["Passed", "Failed"]:
                ws_prj.cell(r, 7).value = "22/10/2026"
            else:
                ws_prj.cell(r, 7).value = None
            ws_prj.cell(r, 7).font = font_tahoma_10
            ws_prj.cell(r, 7).alignment = align_left

    print(f"Project Management updated: {prj_stats}")

    # ==========================================
    # 3. Verify Formulas in Test Statistics & Cover
    # ==========================================
    ws_stat = wb["Test Statistics"]
    assert ws_stat["D12"].value == "='Project Management'!B6", "Formula D12 mismatch"
    assert ws_stat["E12"].value == "='Project Management'!C6", "Formula E12 mismatch"
    assert ws_stat["F12"].value == "='Project Management'!D6", "Formula F12 mismatch"

    assert ws_stat["D14"].value == "='Version Control'!B6", "Formula D14 mismatch"
    assert ws_stat["E14"].value == "='Version Control'!C6", "Formula E14 mismatch"
    assert ws_stat["F14"].value == "='Version Control'!D6", "Formula F14 mismatch"
    print("Formulas in Test Statistics verified 100% intact!")

    # ==========================================
    # 4. Save to Local and OneDrive Targets
    # ==========================================
    print(f"Saving to local: {LOCAL_TARGET} ...")
    wb.save(LOCAL_TARGET)

    print(f"Saving to OneDrive: {ONEDRIVE_TARGET} ...")
    wb.save(ONEDRIVE_TARGET)

    print("Both local and OneDrive workbooks saved successfully!")

if __name__ == "__main__":
    update_test_report()
